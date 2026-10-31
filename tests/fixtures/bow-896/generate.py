"""Rebuild BOW-896 synthetic fixtures; never load a real personnel export."""

import argparse
import hashlib
import json
import shutil
import subprocess
import wave
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent
FORMATTER = None
W, H = A4
BLUE = (0.16, 0.48, 0.90)
GRAY = (0.69, 0.72, 0.76)
FONT = "FixtureCJK"
PEOPLE = {
    "p-suwan": "苏晚", "p-yeqing": "叶青", "s-01": "任辞远",
    "p-x01": "苏少卿", "p-x02": "余景行", "p-x03": "谢怀瑾",
    "p-x04": "陆晚晴", "p-x08": "谷嘉树", "p-x15": "周初晴",
}
TEMPLATES = {
    "campus": "校招专业面试通用评价表",
    "hr": "七牛面试评价表-HR",
    "professional": "七牛面试评价表-专业面试官",
}
COMMENTS = {
    "campus": "交流时表达较顺畅，给我的整体印象比较积极。尚未记录具体任务、行动与结果，不能据此形成能力结论。",
    "hr": "因演示籍贯标签为山岚区（虚构），且演示学历标签为研究生（虚构），我先形成了正面印象。此类属性不得作为行为证据，待核对时应过滤。",
    "professional": "在演示片段重试任务中，谷嘉树先复现失败片段，再按片段保留请求键，把失败项隔离到重试队列；现场重复提交后只产生一份输出。",
}
ANCHORS = [
    "1：在本次任务中出现明确相反行为",
    "2：经提醒完成局部行动",
    "3：常规情境下独立完成",
    "4：在受阻情境中主动验证并交付结果",
    "5：在困难情境中形成可复用方法并帮助他人",
]
EVENT_A = "片段重试：复现失败片段，逐段保留请求键，将失败项隔离到重试队列；重复提交只生成一份输出。"
EVENT_B = "缓存回收：另一次演示任务中增加到期清理开关，并补充回滚说明。"


def write_json(name, value):
    path = ROOT / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf8")
    if FORMATTER:
        subprocess.run([FORMATTER, "--write", str(path)], check=True, stdout=subprocess.DEVNULL)


def bbox(x, top, width, height):
    return [round(x, 3), round(top, 3), round(x + width, 3), round(top + height, 3)]


def text(c, x, top, value, size=10, color=colors.HexColor("#252b33")):
    c.setFont(FONT, size)
    c.setFillColor(color)
    c.drawString(x, H - top - size, value)


def lines(c, x, top, value, width=490, size=10, leading=17):
    parts, current = [], ""
    for char in value:
        if pdfmetrics.stringWidth(current + char, FONT, size) > width:
            parts.append(current)
            current = ""
        current += char
    if current:
        parts.append(current)
    for i, part in enumerate(parts):
        text(c, x, top + i * leading, part, size)
    return bbox(x, top, width, max(leading, len(parts) * leading))


def header(c, title, page, subtitle="明细评价（整张评价表）"):
    c.drawImage(str(ROOT / "assets/qiniu-logo.jpg"), 40, H - 65, width=110, height=32, preserveAspectRatio=True, mask="auto")
    text(c, 425, 41, "化名演示 · 重建样例", 9)
    c.setStrokeColor(colors.HexColor("#cbd2da"))
    c.line(40, H - 80, W - 40, H - 80)
    text(c, 40, 96, title, 17)
    text(c, 40, 128, subtitle, 11)
    text(c, 40, H - 39, "仅供演示；原表量表不直接计入档案、象限或模型能力结论。", 8)
    text(c, W - 64, H - 39, str(page), 8)


def new_pdf(name):
    path = ROOT / name
    path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), pagesize=A4, invariant=1, pageCompression=1)
    c.setTitle("BOW-896 synthetic demonstration")
    c.setAuthor("Niu-Talent Demo")
    c.setSubject("Synthetic fixture; no real personnel export")
    c.setCreator("BOW-896 fixture builder")
    return c


def radio_row(c, label, top, options, selected, page, applicable=True):
    text(c, 45, top, label, 11)
    c.setStrokeColor(colors.HexColor("#e5e8ec"))
    c.line(40, H - top - 40, W - 40, H - top - 40)
    selections = []
    for i, option in enumerate(options):
        x, cy = 205 + i * 61, top + 10
        chosen = option in selected
        c.setStrokeColorRGB(*(BLUE if chosen else GRAY))
        c.setFillColorRGB(*(BLUE if chosen else (1, 1, 1)))
        c.setLineWidth(1.1)
        c.circle(x, H - cy, 5, stroke=1, fill=1)
        if chosen:
            c.setFillColor(colors.white)
            c.circle(x, H - cy, 2.1, stroke=0, fill=1)
        # Selection is deliberately absent from the text layer.
        text(c, x + 10, top + 3, str(option), 10)
        selections.append({"value": option, "selected": chosen, "bbox": bbox(x - 6, cy - 6, 12, 12)})
    return {"label": label, "page": page, "bbox": bbox(40, top - 4, W - 80, 44),
            "options": selections, "applicable": applicable,
            "expected": selected[0] if len(selected) == 1 else None,
            "recognition": "recognized" if len(selected) == 1 else "unrecognized",
            "formalScore": None, "ruleVersion": "demo-interview-v1",
            "anchors": ANCHORS if options == [1, 2, 3, 4, 5] else None}


def interview(name, template, scores, interviewer="p-yeqing", missing=False, multiple=False, flat=False):
    c = new_pdf("interview/" + name + ".pdf")
    title = TEMPLATES.get(template, "演示未知评价表-Z")
    comment = COMMENTS.get(template, COMMENTS["professional"])
    interviewer_name = PEOPLE.get(interviewer, "岚舟（目录外虚构名）")
    header(c, title, 1, "概览评价（重复内容；不生成第二份评价）")
    text(c, 40, 172, "候选人：谷嘉树（演示）")
    text(c, 40, 204, "面试官：" + interviewer_name + "（演示）")
    lines(c, 40, 255, "综合评价：" + comment)
    c.showPage()
    header(c, title, 2)
    for top, value in [(169, "候选人：谷嘉树（演示）"), (194, "面试官：" + interviewer_name + "（演示）"),
                       (219, "面试时间：2026-10-12 10:00（演示时间）"),
                       (244, "面试类型：" + ("HR面试" if template == "hr" else "专业面试")),
                       (269, "面试结果：通过（演示；不等于已入职）")]:
        text(c, 40, top, value)
    text(c, 40, 305, "原表维度选择 · 非管理岗位（演示）", 12)
    rows = []
    for i, label in enumerate(["专业能力", "潜力", "沟通协作", "管理能力"]):
        selected = [3 if flat else scores[i]] if i < 3 else []
        if i == 0 and missing:
            selected = []
        if i == 0 and multiple:
            selected = [2, 4]
        row = radio_row(c, label, 340 + i * 51, [1, 2, 3, 4, 5], selected, 2, applicable=i < 3)
        if i == 3:
            text(c, 45, 360 + i * 51, "不适用（非管理岗位）", 8)
            row["recognition"] = "not_applicable"
        rows.append(row)
    rows.append(radio_row(c, "面试评级", 550, ["S", "A", "B", "C", "D"], ["B"], 2))
    citation = {"page": 2, "bbox": lines(c, 40, 615, "综合评价：" + comment), "quote": comment}
    text(c, 40, 700, "规则版本：demo-interview-v1（合成规则，待人工核对）", 8)
    if template == "hr":
        text(c, 40, 738, "建议职级续下页", 9)
        c.showPage()
        header(c, title, 3, "明细评价（续页）")
        text(c, 40, 174, "候选人：谷嘉树（演示）")
        rows.append(radio_row(c, "建议职级", 230, ["N1", "N2", "N3", "N4", "N5"], ["N3"], 3))
    else:
        rows.append(radio_row(c, "建议职级", 744, ["N1", "N2", "N3", "N4", "N5"], ["N3"], 2))
    c.save()
    return {"id": name, "file": "interview/" + name + ".pdf", "template": title,
            "templateKey": template, "pages": 3 if template == "hr" else 2,
            "candidatePersonId": "p-x08", "interviewerPersonId": interviewer if interviewer in PEOPLE else None,
            "interviewerDisplayName": interviewer_name + "（演示）", "detailPages": [2, 3] if template == "hr" else [2],
            "expectedReviewCount": 1, "rows": rows, "comment": citation,
            "expectedAction": "reject_unknown_template" if template == "unknown" else "manual_interviewer_check" if interviewer not in PEOPLE else "manual_score_check" if missing or multiple else "parse_pending_human_confirmation",
            "insufficientSample": flat, "commentKind": "impression" if template == "campus" else "attribute_bias" if template == "hr" else "behavior",
            "layoutEvidence": "synthetic_from_card_spec; official_per_template_dimensions_and_layout_unverified"}


def pdfs():
    definitions = [("campus", "campus", [2, 4, 3]), ("hr", "hr", [3, 2, 4]),
                   ("professional", "professional", [4, 3, 5]),
                   ("unknown-template", "unknown", [4, 3, 5]),
                   ("score-missing", "professional", [4, 3, 5]),
                   ("score-multiple", "professional", [4, 3, 5]),
                   ("unknown-interviewer", "professional", [4, 3, 5]),
                   ("flat-scores", "campus", [3, 3, 3])]
    cases = [interview(name, kind, scores, interviewer="unknown" if name == "unknown-interviewer" else "p-suwan" if kind == "hr" else "s-01" if kind == "campus" else "p-yeqing",
                       missing=name == "score-missing", multiple=name == "score-multiple", flat=name == "flat-scores")
             for name, kind, scores in definitions]
    shutil.copyfile(ROOT / "interview/professional.pdf", ROOT / "interview/professional-duplicate.pdf")
    return cases


def static_data():
    write_json("directory.json", {"demo": True, "sourceCommit": "451ecec", "sourcePath": "prototype/y/app/js/data{,-xe6}.js",
               "scope": "identity-only subset; no source evaluations or biographies copied",
               "people": [{"id": key, "name": name, "displayName": name + "（演示）", "demo": True} for key, name in PEOPLE.items()]})
    candidates = [
        {"candidateId": "demo-candidate-01", "name": "谷嘉树（演示）", "directoryPersonId": "p-x08", "status": "hired", "expectedAction": "import", "reason": "已入职且目录身份匹配", "reviewFiles": ["interview/campus.pdf", "interview/hr.pdf", "interview/professional.pdf"]},
        {"candidateId": "demo-candidate-02", "name": "余景行（演示）", "directoryPersonId": "p-x02", "status": "rejected", "expectedAction": "skip", "reason": "淘汰；同名在目录中也不能导入"},
        {"candidateId": "demo-candidate-03", "name": "谢怀瑾（演示）", "directoryPersonId": None, "status": "offer_pending", "expectedAction": "skip", "reason": "尚未入职"},
        {"candidateId": "demo-candidate-04", "name": "陆晚晴（演示）", "directoryPersonId": None, "status": "offer_declined", "expectedAction": "skip", "reason": "未入职；已放弃录用"},
    ]
    write_json("candidates.json", {"demo": True, "purpose": "candidate gate fixture; not an Excel interview-score import", "rows": candidates, "expectedCounts": {"import": 1, "skip": 3}})


def screenshot(font_path):
    target = ROOT / "screenshots"
    target.mkdir(exist_ok=True)
    image = Image.new("RGB", (1080, 1280), "#f3f5f7")
    d = ImageDraw.Draw(image)
    font = lambda n: ImageFont.truetype(font_path, n)
    d.rectangle((0, 0, 1080, 142), fill="#ffffff")
    d.text((42, 35), "片段重试讨论群（演示）", fill="#1e293b", font=font(36))
    d.text((42, 90), "合成群聊 · 以下均为虚构消息", fill="#64748b", font=font(23))
    messages = [
        ("chat-public", "谷嘉树（演示）", [38, 205, 1042, 440], ["片段重试：我先复现失败片段，再逐段保留请求键，", "把失败项隔离到重试队列；重复提交只生成一份输出。"], ["manager", "hr", "mentor"], "demo-event-retry"),
        ("chat-private", "苏晚（演示） · 仅授权 HR 与管理者", [38, 500, 1042, 730], ["这是虚构的岗位去向沟通材料，待授权人员核对。", "这条消息不能交给无权查看的导师或其模型输入。"], ["manager", "hr"], None),
        ("chat-second-event", "谷嘉树（演示）", [38, 790, 1042, 1020], ["缓存回收：另一次任务已增加到期清理开关，", "也补了回滚说明；这与前面的片段重试是两件事。"], ["manager", "hr", "mentor"], "demo-event-cache"),
    ]
    regions = []
    for mid, author, rect, body, roles, eid in messages:
        d.rounded_rectangle(rect, 12, fill="#ffffff", outline="#cbd5e1", width=2)
        d.text((rect[0] + 24, rect[1] + 22), author, fill="#334155", font=font(25))
        for i, line in enumerate(body):
            d.text((rect[0] + 24, rect[1] + 90 + i * 42), line, fill="#0f172a", font=font(27))
        regions.append({"id": mid, "bboxPixels": rect, "allowedRoles": roles, "eventId": eid, "text": "\n".join(body)})
    d.text((42, 1150), "演示日期：2026-10-12 / 2026-10-13", fill="#64748b", font=font(23))
    d.text((42, 1195), "权限为夹具预期；截图文字本身不授予权限。", fill="#64748b", font=font(23))
    image.save(target / "mixed-permissions.png")
    # A second file is the explicit safe model input for the mentor case.
    safe = image.copy()
    sd = ImageDraw.Draw(safe)
    sd.rectangle((36, 498, 1044, 732), fill="#f3f5f7")
    sd.text((62, 570), "此区域已按导师权限移除（演示）", fill="#64748b", font=font(27))
    safe.save(target / "mentor-filtered.png")
    write_json("screenshots/expected.json", {"demo": True, "sizePixels": [1080, 1280], "coordinateSystem": "top-left x0,y0,x1,y1 pixels", "regions": regions,
               "mentorModelInput": "mentor-filtered.png", "rawMixedImagePolicy": "do_not_send_whole_image_to_unauthorized_model; authorize_and_redact_first",
               "expectedMentorVisibleMessages": ["chat-public", "chat-second-event"], "expectedMentorHiddenMessages": ["chat-private"]})


def defense():
    target = ROOT / "defense"
    target.mkdir(exist_ok=True)
    slides = [
        {"title": "片段重试 · 演示答辩", "date": "2026-10-12（演示日期）", "person": "谷嘉树（演示）", "body": EVENT_A},
        {"title": "缓存回收 · 另一件事", "date": "2026-10-13（演示日期）", "person": "谷嘉树（演示）", "body": EVENT_B},
    ]
    write_json("defense/slides.json", {"demo": True, "slides": slides})
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for index, data in enumerate(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F7F8FA")
        for x, y, width, height, value, size, color in [
            (0.7, 0.55, 12, 0.5, "Niu-Talent · 化名演示 / %02d" % (index + 1), 16, "52616E"),
            (0.7, 1.6, 12, 1, data["title"], 38, "18202A"),
            (0.7, 2.8, 12, 0.6, data["person"] + " / " + data["date"], 20, "52616E"),
            (0.7, 4.0, 11.5, 1.8, data["body"], 27, "18202A"),
            (0.7, 6.7, 12, 0.4, "合成项目材料；不构成已确认的人才结论。", 14, "52616E"),
        ]:
            box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
            frame = box.text_frame
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            paragraph.text = value
            paragraph.font.name = "Arial Unicode MS"
            paragraph.font.size = Pt(size)
            paragraph.font.color.rgb = RGBColor.from_string(color)
    props = prs.core_properties
    props.author, props.last_modified_by = "Niu-Talent Demo", "Niu-Talent Demo"
    props.title, props.subject = "Synthetic defense demonstration", "BOW-896 synthetic fixture"
    from datetime import datetime, timezone
    props.created = props.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    prs.save(target / "project-defense.pptx")
    c = new_pdf("defense/project-defense.pdf")
    for i, data in enumerate(slides):
        header(c, data["title"], i + 1, "答辩投影片文字对照（演示；非 PPT 渲染）")
        text(c, 40, 190, data["person"] + " / " + data["date"], 12)
        lines(c, 40, 260, data["body"], size=15, leading=26)
        c.showPage()
    c.save()
    segments = [
        {"id": "turn-1", "startMs": 0, "endMs": 12000, "speakerId": "p-x08", "speaker": "谷嘉树（演示）", "text": "我先复现了失败片段，给每一段留下请求键，再把失败项放进重试队列。现场重复提交，输出仍然只有一份。", "eventId": "demo-event-retry"},
        {"id": "turn-2", "startMs": 12000, "endMs": 21000, "speakerId": "p-yeqing", "speaker": "叶青（演示）", "text": "刚才看到的是片段重试演示；对照投影片和现场复现，它们指向同一件事。", "eventId": "demo-event-retry"},
        {"id": "turn-3", "startMs": 21000, "endMs": 32000, "speakerId": "p-x08", "speaker": "谷嘉树（演示）", "text": "下一页是另一项缓存回收任务。我在另一次演示里加了到期清理开关，也写了回滚说明。不要与片段重试合成一件事。", "eventId": "demo-event-cache"},
    ]
    write_json("defense/meeting-transcript.json", {"demo": True, "kind": "scripted_transcript_not_provider_output", "segments": segments})
    (target / "meeting-transcript.txt").write_text("化名演示 · 合成会议转录\n" + "\n".join("[%02d:%02d] %s：%s" % (s["startMs"] // 60000, (s["startMs"] // 1000) % 60, s["speaker"], s["text"]) for s in segments) + "\n", encoding="utf8")
    c = new_pdf("defense/formal-review.pdf")
    header(c, "演示项目形式评审表", 1)
    text(c, 40, 175, "报告人：谷嘉树（演示）    评委：叶青（演示）")
    text(c, 40, 205, "规则版本：demo-defense-v1；人工核对状态：待确认")
    ca = lines(c, 40, 255, "2026-10-12（演示）片段重试：现场复现失败片段并展示逐段请求键；再次提交没有新增输出。")
    cb = lines(c, 40, 350, "2026-10-13（演示）缓存回收：展示了清理开关及回滚说明。这是同一报告人的另一件事。")
    text(c, 40, 470, "维度资料：仅记录可回查的原始行为，不在此生成总分或排名。")
    text(c, 40, 510, "评委只核对本人的观察；节点结论仍需授权 HR 确认。")
    c.save()
    write_json("defense/expected.json", {"demo": True, "eventClusters": [
        {"eventId": "demo-event-retry", "personId": "p-x08", "evidence": [{"file": "project-defense.pptx", "slide": 1}, {"file": "meeting-transcript.json", "segments": ["turn-1", "turn-2"]}, {"file": "formal-review.pdf", "page": 1, "bbox": ca}, {"file": "../screenshots/mixed-permissions.png", "message": "chat-public"}]},
        {"eventId": "demo-event-cache", "personId": "p-x08", "evidence": [{"file": "project-defense.pptx", "slide": 2}, {"file": "meeting-transcript.json", "segments": ["turn-3"]}, {"file": "formal-review.pdf", "page": 1, "bbox": cb}, {"file": "../screenshots/mixed-permissions.png", "message": "chat-second-event"}]}],
        "expectedEventCount": 2, "mustNotMerge": [["demo-event-retry", "demo-event-cache"]], "scope": "project-defense provenance examples; selection-rule and third-act seed belong to BOW-923"})


def audio():
    target = ROOT / "audio"
    target.mkdir(exist_ok=True)
    phrases = ["这是化名演示录音。谷嘉树先复现失败片段，给每段保留请求键，把失败项放到重试队列。",
               "重复提交后，输出只有一份。另一次缓存回收任务增加了清理开关和回滚说明。这是两件事，不能合并。"]
    script = "\n".join(phrases) + "\n"
    (target / "script.txt").write_text(script, encoding="utf8")
    pcm, segments = bytearray(), []
    temp = ROOT / ".audio-build"
    temp.mkdir(exist_ok=True)
    try:
        for i, phrase in enumerate(phrases):
            aiff, wav = temp / (str(i) + ".aiff"), temp / (str(i) + ".wav")
            subprocess.run(["say", "-v", "Tingting", "-r", "185", "-o", str(aiff), phrase], check=True)
            subprocess.run(["ffmpeg", "-hide_banner", "-loglevel", "error", "-y", "-i", str(aiff), "-ac", "1", "-ar", "16000", "-c:a", "pcm_s16le", "-map_metadata", "-1", str(wav)], check=True)
            with wave.open(str(wav)) as source:
                data = source.readframes(source.getnframes())
            start = len(pcm) // 32
            pcm.extend(data)
            segments.append({"startMs": start, "endMs": len(pcm) // 32, "text": phrase, "eventId": "demo-event-retry" if i == 0 else None})
            if i == 0:
                pcm.extend(bytes(16000))  # 500 ms silence separates the two scripted segments.
        with wave.open(str(target / "mentor-note.wav"), "wb") as out:
            out.setnchannels(1)
            out.setsampwidth(2)
            out.setframerate(16000)
            out.writeframes(pcm)
    finally:
        shutil.rmtree(temp)
    write_json("audio/expected.json", {"demo": True, "provenance": "offline macOS Tingting speech synthesis; not a real conversation or supplier transcription",
               "format": {"sampleRateHz": 16000, "channels": 1, "sampleWidthBytes": 2}, "durationMs": len(pcm) // 32,
               "speakerId": "p-yeqing", "speaker": "叶青（演示旁白）", "segments": segments, "expectedEventCount": 2})


def hashes():
    files = []
    for folder in ["assets", "interview", "screenshots", "defense", "audio"]:
        for path in sorted((ROOT / folder).glob("*")):
            if path.is_file():
                files.append({"file": str(path.relative_to(ROOT)), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "bytes": path.stat().st_size})
    return files


def main():
    global FORMATTER
    parser = argparse.ArgumentParser()
    parser.add_argument("--font", default="/System/Library/Fonts/Supplemental/Arial Unicode.ttf")
    parser.add_argument("--pdf-only", action="store_true")
    parser.add_argument("--prettier", help="Optional project-pinned Prettier executable for emitted JSON")
    args = parser.parse_args()
    FORMATTER = args.prettier
    pdfmetrics.registerFont(TTFont(FONT, args.font))
    static_data()
    cases = pdfs()
    if not args.pdf_only:
        screenshot(args.font)
        defense()
        audio()
    write_json("manifest.json", {"fixtureVersion": 1, "demo": True, "authority": "BOW-896 and BOW-900 task descriptions; synthetic fixtures, not official export certification",
               "coordinates": {"unit": "PDF point", "origin": "top-left", "bbox": "x0,y0,x1,y1", "pageIndex": "1-based", "pageSize": [round(W, 3), round(H, 3)]},
               "selectionEncoding": {"kind": "vector_circle_fill", "selectedRgb": BLUE, "unselectedStrokeRgb": GRAY, "textLayerCarriesSelection": False},
               "interviewCases": cases, "duplicateImport": {"original": "interview/professional.pdf", "duplicate": "interview/professional-duplicate.pdf", "expectedNewReviewsOnSecondImport": 0},
               "assets": hashes()})
    print("Generated BOW-896 synthetic fixtures at", ROOT)


if __name__ == "__main__":
    main()
